# -*- coding: utf-8 -*-
"""Output 10 - September 2026 round: Kentron counts as 53 zones, the survey as 203.

Five Kentron zones (K02-K05, K18; 52 spaces) carry no occupancy records - parking was
not available there on the survey day (debrief: Moskovyan made one-way). At the
user's direction (29 Sep 2026) they are not counted at all: Kentron is 53 zones and
the survey 203 zones, which is also the figure the pricing paragraph already uses.

Applies to BOTH Output 10 files - the clean 28092026 revision and the Logit response
copy - and rewords the C61 reply in the response copy.

Run:  python apply_o10_sep_kentron53.py        (edits in place; files must be closed)
"""
import sys
import zipfile
import io

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from apply_o10_sep_comments import _clone_run_after
from docx_edit import RED, find

SEP = "C:/Users/user/Yerevan-Parking/Final Presentation/SEP'26/"
FILES = [SEP + "Output 10 - Parking Analysis Report - 28092026.docx",
         SEP + "Output 10_Parking Analysis Report_25082026_Logit comments - STS responses.docx"]

OLD_COUNTS = ("(Kentron 58, of which the 53 swept for occupancy are the basis of every Kentron "
              "occupancy figure in this report; Komitas 53;")
NEW_COUNTS = "(Kentron 53; Komitas 53;"
OLD_REPLY = ("Both correct: Kentron has 58 zones (53 surveyed), Komitas 53. Added a definition "
             "of zone and the counts per area.")
NEW_REPLY = "Both correct: Kentron 53, Komitas 53. Added a definition of zone and the counts per area."


def fix_doc(path):
    doc = Document(path)
    done = maybe_done(doc)
    if done is not None:
        print("already corrected:", path.split("/")[-1])
        return done
    par = find(doc, "10,135 distinct vehicles across 208 survey zones")

    # 208 -> red 203, keeping the surrounding black text black
    host = next(r for r in par.runs if "across 208 survey zones" in r.text)
    head, tail = host.text.split("across 208 survey zones", 1)
    host.text = head + "across "
    red = _clone_run_after(host, "203", red=True)
    rest = _clone_run_after(red, " survey zones" + tail, red=False)
    rest.font.color.rgb = host.font.color.rgb

    counts = [r for r in par.runs if OLD_COUNTS in r.text]
    assert len(counts) == 1, "zone-count parenthetical not found in a single run"
    counts[0].text = counts[0].text.replace(OLD_COUNTS, NEW_COUNTS)

    # drop the three red " surveyed" insertions
    dropped = 0
    for p in doc.paragraphs:
        for r in list(p.runs):
            if r.text == " surveyed":
                r._r.getparent().remove(r._r)
                dropped += 1
    assert dropped == 3, f"expected 3 ' surveyed' runs, found {dropped}"

    body = "\n".join(p.text for p in doc.paragraphs)
    for gone in ("Kentron 58", "208 survey zones", "53 surveyed zones", "53 swept"):
        assert gone not in body, gone
    assert "across 203 survey zones - short segments of street" in body
    doc.save(path)
    return par.text


def maybe_done(doc):
    """Paragraph text if this file was already corrected (so a re-run after a locked
    file only finishes the rest), else None."""
    body = "\n".join(p.text for p in doc.paragraphs)
    if "across 203 survey zones - short segments of street" in body and "Kentron 58" not in body:
        return find(doc, "10,135 distinct vehicles across 203 survey zones").text
    return None


def fix_reply(path):
    zin = zipfile.ZipFile(path)
    parts = {n: zin.read(n) for n in zin.namelist()}
    zin.close()
    xml = parts["word/comments.xml"].decode("utf-8")
    if NEW_REPLY in xml:
        print("reply already updated")
        return
    assert xml.count(OLD_REPLY) == 1, "C61 reply not found"
    parts["word/comments.xml"] = xml.replace(OLD_REPLY, NEW_REPLY).encode("utf-8")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in parts.items():
            zout.writestr(name, data)
    with open(path, "wb") as fh:
        fh.write(buf.getvalue())


def main():
    texts = [fix_doc(f) for f in FILES]
    assert texts[0] == texts[1], "the two files diverged"
    fix_reply(FILES[1])
    i = texts[0].find("across 203")
    print(texts[0][i - 20:i + 320])


if __name__ == "__main__":
    main()
