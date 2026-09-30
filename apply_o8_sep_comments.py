# -*- coding: utf-8 -*-
"""Output 8 - Logit comments of 1/24 Sep 2026: STS response copy, pass 1 (text edits).

Works on a COPY of the reviewers' commented file (the Logit original is never touched),
same workflow as Output 10 (build_o10_logit_responses.py). New text is red; deleted text
is simply removed. Comment anchors (commentRangeStart/End + reference runs) are kept.

  C79  occupancy explanation rewritten shorter, with the Kentron worked example
  C91  "Measured Displacement by Area" opening rewritten: 1,756 -> 1,123 -> 1,886
  C66  Displacement Assessment chapter moved after the Field Occupancy chapter, with
       "Measured Displacement by Area" folded into it; the duplicate "Locally measured"
       paragraph is dropped and its caveat carried into the C91 paragraph
  C97  conclusion caveat: most off-street stock is gated courtyards, not public parking

Captions are SEQ fields, so renumbering happens on a field update in Word (COM).

Run:  python apply_o8_sep_comments.py        (refuses to overwrite the output)
"""
import os

from docx import Document
from docx.oxml.ns import qn

from apply_o10_sep_comments import _clone_run_after
from docx_edit import RED, find

SEP = "C:/Users/user/Yerevan-Parking/Final Presentation/SEP'26/"
SRC = SEP + "Output 8 - Parking Surveys and Analysis Report_25082026_Logit comments.docx"
DST = SEP + "Output 8 - Parking Surveys and Analysis Report_25082026_Logit comments - STS responses.docx"

C79_OCCUPANCY = (
    "How occupancy is measured. Occupancy is the number of distinct vehicles that used "
    "a stretch of kerb during an hour, divided by the number of spaces on it. Each area "
    "is reported at its busiest hour. In Kentron, for example, 473 vehicles used the 534 "
    "surveyed spaces in the 21:00 hour: 473 / 534 = 89%. Every area percentage in this "
    "chapter is calculated this way. The same calculation for a single zone uses that "
    "zone's own busiest hour, which shows which stretches of kerb are saturated; because "
    "those hours differ from zone to zone, zone figures cannot be added up into an area "
    "figure."
)
C79_OVER_100 = (
    "A zone can read above 100%. The count covers every vehicle that used the kerb "
    "during the hour, not how many stood there at one moment, so three cars using one "
    "space in turn read as 300%; the busiest zone in the survey reaches six times its "
    "capacity. An area could do the same, although none does: the highest is Gai "
    "Avenue at 99%."
)
C91_DISPLACEMENT = (
    "The conceptual design identifies which surveyed zones lose their parking, so "
    "displacement can be measured rather than assumed. At the busiest hour, 1,756 "
    "vehicles used the surveyed kerb (previous chapter), but only those parked in the "
    "removed zones are displaced. The design removes 1,886 spaces in the six areas, and "
    "at the busiest hour for those spaces in each area 1,123 vehicles were parked in "
    "them. That is the displaced demand: 60% of the spaces removed. It does not mean "
    "that about 40% of the removed spaces are unused, only that they are never all "
    "occupied in the same hour."
)
C97_CAVEAT = (
    " Most of the off-street stock (5,601 spaces) is gated residential courtyards that "
    "are not open for public parking today."
)


def rewrite_keep_comments(par, text):
    """Whole-paragraph rewrite in red that keeps comment markers and reference runs
    (docx_edit.set_text / rewrite_red would delete the reference run)."""
    text_runs = [r for r in par.runs if r._r.find(qn("w:t")) is not None
                 and r._r.find(qn("w:commentReference")) is None]
    first = text_runs[0]
    for r in text_runs[1:]:
        r._r.getparent().remove(r._r)
    first.text = text
    first.font.color.rgb = RED
    first.bold = None


def append_red_to_text(par, text):
    runs = [r for r in par.runs if r._r.find(qn("w:t")) is not None and r.text]
    return _clone_run_after(runs[-1], text, red=True)


def heading(doc, text, style):
    hits = [p for p in doc.paragraphs if p.text.strip() == text and p.style.name == style]
    if len(hits) != 1:
        raise LookupError(f"{len(hits)} {style} headings named {text!r}")
    return hits[0]


def block(first, stop):
    """Body elements from `first` up to (not including) `stop`."""
    out, el = [], first._p
    while el is not stop._p:
        if el is None:
            raise LookupError("block: stop element never reached")
        out.append(el)
        el = el.getnext()
    return out


def move_block(elements, before):
    for el in elements:
        el.getparent().remove(el)
        before._p.addprevious(el)


def apply_edits(doc):
    # C79
    rewrite_keep_comments(find(doc, "How occupancy is measured, and at what scale"), C79_OCCUPANCY)
    rewrite_keep_comments(find(doc, "A zone reading above 100% is not a data error"), C79_OVER_100)

    # C91
    rewrite_keep_comments(find(doc, "Because the conceptual design identifies which of the surveyed"),
                          C91_DISPLACEMENT)

    # C97
    append_red_to_text(find(doc, "Available absorptive capacity includes 6,732 off-street"), C97_CAVEAT)

    # C66: the "Locally measured" paragraph duplicates the C91 paragraph once the two meet
    locally = find(doc, "Locally measured")
    locally._p.getparent().remove(locally._p)

    disp = heading(doc, "Displacement Assessment", "Heading 1")
    field = heading(doc, "Field Occupancy Survey: Methodology, Locations and Results", "Heading 1")
    measured = heading(doc, "Measured Displacement by Area", "Heading 2")
    absorb = heading(doc, "Available Absorptive Capacity", "Heading 2")
    concl = heading(doc, "Conclusions", "Heading 1")

    disp_intro = block(disp, absorb)          # H1 + Understanding Displacement
    disp_absorb = block(absorb, field)        # Available Absorptive Capacity (+ trailing blanks)
    measured_blk = block(measured, concl)     # Measured Displacement by Area

    # new order: ... Field Occupancy | Displacement: intro, measured, absorptive | Conclusions
    move_block(disp_intro, concl)
    move_block(measured_blk, concl)
    move_block(disp_absorb, concl)


def check_order(doc):
    heads = [p.text.strip() for p in doc.paragraphs if p.style.name in ("Heading 1", "Heading 2")]
    want = ["Impact by Sensitivity Zone",
            "Field Occupancy Survey: Methodology, Locations and Results",
            "Survey Design and Instrument", "Where and When the Surveys Were Conducted",
            "Occupancy Results by Area", "Duration of Stay and Turnover",
            "Displacement Assessment", "Understanding Displacement",
            "Measured Displacement by Area", "Available Absorptive Capacity", "Conclusions"]
    i = heads.index(want[0])
    if heads[i:i + len(want)] != want:
        raise SystemExit(f"unexpected heading order: {heads[i:]}")


def main():
    if os.path.exists(DST):
        raise SystemExit(f"refusing to overwrite {DST}")
    doc = Document(SRC)
    apply_edits(doc)
    check_order(doc)
    doc.save(DST)
    print("wrote", DST)


if __name__ == "__main__":
    main()
