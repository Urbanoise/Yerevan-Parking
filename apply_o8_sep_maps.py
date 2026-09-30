# -*- coding: utf-8 -*-
"""Output 8 - Logit comments of 1/24 Sep 2026: STS response copy, pass 2 (maps + zone).

Edits the response copy written by apply_o8_sep_comments.py in place (a backup is
taken first). Figures are prints of the web app's static build, captured and
annotated by capture_o8_app_maps.py (outputs in O8 app maps/).

  C28  supply maps: regulatory zone, configuration, signage, marking, and a
       street-level map per corridor with the off-street lots
  C60  retained vs removed on-corridor parking under the conceptual design
  C72  location of the six survey areas
  C66  average daily occupancy by zone, one panel per area
  C69  "Zones?": zone defined at first use as a street segment (as in Output 10)

Caption anchors are matched on the Caption style: an unstyled find() hits the List of
Figures entry first (this put four maps inside the list on the first run, 30 Sep;
fix_o8_sep_map_positions.py moved them). Captions carry SEQ fields; renumber by updating fields in Word (COM) afterwards.

Run:  python apply_o8_sep_maps.py            (dry run, saves nothing)
      python apply_o8_sep_maps.py --apply
"""
import shutil
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

from apply_o8_sep_comments import DST
from apply_o10_sep_comments import insert_red
from docx_edit import add_picture_after, find, find_styled, set_caption

MAPS = "C:/Users/user/Yerevan-Parking/Final Presentation/SEP'26/O8 app maps/"
BACKUP = DST.replace(".docx", " (before maps).docx")
W = 6.3

ZONE_DEF = " - short segments of street, typically one block face -"


def before(anchor):
    """An empty paragraph just before `anchor` (a heading), to hang a figure on."""
    p = OxmlElement("w:p")
    anchor._p.addprevious(p)
    return Paragraph(p, anchor._parent)


def figure(after, png, title, inches=W):
    cap = add_picture_after(after, MAPS + png, inches=inches, caption=title)
    set_caption(cap, "Figure", title)
    return cap


def heading(doc, text):
    hits = [p for p in doc.paragraphs if p.text.strip() == text and p.style.name.startswith("Heading")]
    if len(hits) != 1:
        raise LookupError(f"{len(hits)} headings named {text!r}")
    return hits[0]


def apply_edits(doc):
    # C69
    par = find(doc, "Enumerators walked a fixed circuit of numbered survey zones")
    insert_red(par, "numbered survey zones", ZONE_DEF)

    # C28
    figure(find_styled(doc, "Figure 1 - On-corridor on-street supply by regulatory zone", "Caption"), "reg.jpg",
           "Location of on-corridor on-street parking by regulatory zone")
    figure(find_styled(doc, "Figure 2 - On-corridor on-street supply by marked configuration", "Caption"), "method.jpg",
           "Location of on-corridor on-street parking by configuration")
    figure(find_styled(doc, "Figure 4 - Parking signage coverage", "Caption"), "signage.jpg",
           "Location of parking signage")
    figure(find_styled(doc, "Figure 5 - Pavement marking coverage", "Caption"), "marking.jpg",
           "Location of pavement marking")
    figure(before(heading(doc, "Corridor 2: Street-Level Detail")), "loc_c1.jpg",
           "Corridor 1: on-street parking by location, with off-street lots")
    figure(before(heading(doc, "Parking Impacts of the Conceptual Corridor Designs")), "loc_c2.jpg",
           "Corridor 2: on-street parking by location, with off-street lots")

    # C60
    figure(find_styled(doc, "Figure 6 - Parking removed as a share of the impact area", "Caption"), "removed.jpg",
           "On-corridor parking retained and removed in the conceptual design")

    # C72
    figure(find(doc, "Six areas were surveyed rather than the full length of the corridors"), "areas.jpg",
           "Location of the six field survey areas")

    # C66
    figure(find_styled(doc, "Figure 7 - Peak-hour occupancy by surveyed area", "Caption"), "occ_grid.jpg",
           "Average daily occupancy by zone in the six survey areas", inches=6.8)


def main(apply):
    doc = Document(DST)
    body = "\n".join(p.text for p in doc.paragraphs)
    if "short segments of street" in body:
        raise SystemExit("already applied - refusing to apply twice")
    apply_edits(doc)
    if apply:
        shutil.copyfile(DST, BACKUP)
        doc.save(DST)
        print("saved", DST, "\nbackup", BACKUP)
    else:
        print("dry run OK - nothing saved")


if __name__ == "__main__":
    main("--apply" in sys.argv)
