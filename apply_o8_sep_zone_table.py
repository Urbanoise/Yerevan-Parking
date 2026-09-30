# -*- coding: utf-8 -*-
"""Output 8 - Logit comments, Sep 2026: zone table (C68), clearer C70 reply, Figure 16.

  C68  Livia asked for occupancy results by zone. Per the user (30 Sep): a simple table
       right after the commented paragraph - per area: zones surveyed, spaces, average
       daily occupancy, turnover. Built as a red copy of Table 8 ("areas, dates and
       footprint") so it matches the report's table style. Reply added.
  C70  the reply "Done." was too vague for "The text is very difficult to follow";
       it now says what changed.
  Fig  "Average daily occupancy by zone" re-composed with white gutters between the
       six area panels (compose_o8_app_maps.py); image swapped in place.

Figures come from field-surveys.geojson on the same basis as the rest of the chapter:
203 zones with occupancy data, 2,878 spaces (Kentron 534, not the 586 inventoried).
Average occupancy = capacity-weighted mean of each zone's occupancy_pct (vehicle-hours
over the survey window / capacity); turnover = parking events / spaces, which
reproduces Table 9's turnover column exactly.

Run:  python apply_o8_sep_zone_table.py            (dry run)
      python apply_o8_sep_zone_table.py --apply
"""
import copy
import io
import json
import sys
import zipfile
from collections import defaultdict

from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph
from lxml import etree

import build_o10_logit_responses as o10
from apply_o8_sep_comments import DST
from docx_edit import find, replace_image, set_caption

GEO = "C:/Users/user/Yerevan-Parking/app/static/data/wgs84/field-surveys.geojson"
FIG = "C:/Users/user/Yerevan-Parking/Final Presentation/SEP'26/O8 app maps/occ_grid.jpg"
RED = RGBColor(0xC0, 0x00, 0x00)
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"

AREAS = [("kentron", "Kentron"), ("komitas", "Komitas Avenue"), ("mega", "Gai Avenue (Mega Mall)"),
         ("garegin", "Garegin Nzhdeh"), ("shiraz", "Shiraz / Hasratyan"),
         ("malatia", "Malatia-Sebastia")]
TITLE = "Surveyed zones by area: capacity, average occupancy and turnover"
HEADER = ["Area", "Zones surveyed", "Spaces", "Average daily occupancy", "Turnover per space/day"]
C68_REPLY = "Done, summary table by area added below this paragraph."
C86_REPLY = ("Paragraph rewritten to step from the 1,756 vehicles on all surveyed kerb to the "
             "1,123 parked in the zones the design removes. 1,431 is the pocket-parking "
             "supply (Table 4), not a vehicle count.")
C70_REPLY = "Rewritten shorter and in plain terms, with a worked example (Kentron at 21:00)."


def rows():
    feats = json.load(open(GEO, encoding="utf-8"))["features"]
    by = defaultdict(list)
    for f in feats:
        p = f["properties"]
        if p.get("peak_occupancy") is not None:
            by[p["area"]].append(p)

    def line(label, zs):
        sp = sum(p["space"] for p in zs)
        occ = sum(p["occupancy_pct"] * p["space"] for p in zs) / sp
        turn = sum(p["parking_events"] for p in zs) / sp
        return [label, str(len(zs)), f"{sp:,}", f"{occ:.0f}%", f"{turn:.1f}"]

    out = [HEADER] + [line(lbl, by[k]) for k, lbl in AREAS]
    out.append(line("All six areas", [p for k, _ in AREAS for p in by[k]]))
    assert out[1][1:3] == ["53", "534"] and out[-1][1:3] == ["203", "2,878"], out
    assert [r[4] for r in out[1:7]] == ["7.5", "5.5", "6.6", "5.3", "3.7", "3.5"], out  # = Table 9
    return out


def set_cell_red(cell, text):
    p = cell.paragraphs[0]
    runs = p.runs
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    run = runs[0] if runs else p.add_run()
    run.text = text
    run.font.color.rgb = RED


def add_table(doc):
    anchor = find(doc, "Occupancy is measured against the capacity already recorded")
    model = next(t for t in doc.tables if t.rows[0].cells[1].text.strip() == "Survey date")
    model_cap = Paragraph(model._tbl.getprevious(), anchor._parent)
    if not model_cap.text.startswith("Table 8 - Field occupancy survey"):
        raise SystemExit("Table 8 caption not found above its table")

    cap = copy.deepcopy(model_cap._p)
    tbl = copy.deepcopy(model._tbl)
    for el in (cap, tbl):  # clones must not reuse paragraph ids or bookmarks
        for p in el.iter(qn("w:p")):
            p.attrib.pop(f"{{{W14}}}paraId", None)
            p.attrib.pop(f"{{{W14}}}textId", None)
        for b in list(el.iter(qn("w:bookmarkStart"))) + list(el.iter(qn("w:bookmarkEnd"))):
            b.getparent().remove(b)
    anchor._p.addnext(cap)
    cap.addnext(tbl)
    set_caption(Paragraph(cap, anchor._parent), "Table", TITLE)

    from docx.table import Table
    table = Table(tbl, anchor._parent)
    data = rows()
    if len(table.rows) != len(data):
        raise SystemExit("row count mismatch with the Table 8 template")
    for r, values in zip(table.rows, data):
        for c, v in zip(r.cells, values):
            set_cell_red(c, v)
    return data


def swap_figure(doc):
    cap = next(p for p in doc.paragraphs if p.style.name == "Caption"
               and p.text.endswith("Average daily occupancy by zone in the six survey areas"))
    pic = Paragraph(cap._p.getprevious(), cap._parent)
    replace_image(doc, pic, FIG)
    from PIL import Image
    w, h = Image.open(FIG).size
    for tag in ("wp:extent", "a:ext"):
        for e in pic._p.iter(qn(tag)):
            if e.get("cx"):
                e.set("cy", str(round(int(e.get("cx")) * h / w)))


def comment_id(path, author, starts):
    """Word renumbers comment ids on every save, so find the comment by its text."""
    com = etree.fromstring(zipfile.ZipFile(path).read("word/comments.xml"))
    hits = [c.get(f"{{{o10.W}}}id") for c in com.findall(f"{{{o10.W}}}comment")
            if c.get(f"{{{o10.W}}}author") == author
            and "".join(t.text or "" for t in c.iter(f"{{{o10.W}}}t")).startswith(starts)]
    if len(hits) != 1:
        raise SystemExit(f"{len(hits)} comments by {author} starting {starts!r}")
    return hits[0]


def fix_c70_reply(path):
    set_reply(path, "Lucas Melo", "The text is very difficult to follow", C70_REPLY)


def set_reply(path, author, starts, text):
    """Rewrite the text of Nikoloz's reply under the comment by `author` starting `starts`."""
    cid = comment_id(path, author, starts)
    z = zipfile.ZipFile(path)
    parts = {n: z.read(n) for n in z.namelist()}
    z.close()
    W, W15 = o10.W, o10.W15
    com, ext = etree.fromstring(parts["word/comments.xml"]), etree.fromstring(parts["word/commentsExtended.xml"])
    ns = {"w": W, "w14": W14, "w15": W15}
    parent = com.xpath(f"w:comment[@w:id='{cid}']", namespaces=ns)[0]
    ppid = parent.findall(f"{{{W}}}p")[-1].get(f"{{{W14}}}paraId")
    kids = {e.get(f"{{{W15}}}paraId") for e in ext.xpath(f"w15:commentEx[@w15:paraIdParent='{ppid}']", namespaces=ns)}
    replies = [c for c in com.xpath(f"w:comment[@w:author='{o10.AUTHOR}']", namespaces=ns)
               if c.find(f"{{{W}}}p").get(f"{{{W14}}}paraId") in kids]
    if len(replies) != 1:
        raise SystemExit(f"{len(replies)} replies under comment {cid}")
    ts = list(replies[0].iter(f"{{{W}}}t"))
    ts[0].text = text
    for t in ts[1:]:
        t.text = ""
    parts["word/comments.xml"] = etree.tostring(com, xml_declaration=True, encoding="UTF-8", standalone=True)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in parts.items():
            zout.writestr(name, data)
    open(path, "wb").write(buf.getvalue())


def fix_table_xref(path):
    """The new table is numbered before "areas, dates and footprint", which becomes
    Table 9; the one hard-coded reference to it ("spaces in Table 8") must follow.
    Only the digit changes, in red."""
    doc = Document(path)
    par = find(doc, "fewer than the 2,930 spaces in Table 8")
    run = next((r for r in par.runs if "in Table 8" in r.text), None)
    if run is None:
        raise SystemExit("'in Table 8' straddles runs")
    before, after = run.text.split("in Table 8", 1)
    run.text = before + "in Table "
    red = copy.deepcopy(run._r)
    run._r.addnext(red)
    from docx.text.run import Run
    red_run = Run(red, par)
    red_run.text = "9"
    red_run.font.color.rgb = RED
    tail = copy.deepcopy(run._r)
    red.addnext(tail)
    Run(tail, par).text = after
    assert "spaces in Table 9, because" in par.text
    doc.save(path)


def zones_203(path):
    """User, 30 Sep: 208 -> 203 survey zones, as in Output 10 (five of Kentron's 58
    inventoried zones were not swept). Only the number changes, in red."""
    doc = Document(path)
    par = find(doc, "10,135 distinct vehicles across 208 survey zones")
    run = next((r for r in par.runs if "208" in r.text), None)
    if run is None:
        raise SystemExit("'208' not found in a single run")
    before, after = run.text.split("208", 1)
    run.text = before
    red = copy.deepcopy(run._r)
    run._r.addnext(red)
    from docx.text.run import Run
    Run(red, par).text = "203"
    Run(red, par).font.color.rgb = RED
    tail = copy.deepcopy(run._r)
    red.addnext(tail)
    Run(tail, par).text = after
    assert "across 203 survey zones" in par.text
    doc.save(path)


def replies(path):
    fix_c70_reply(path)
    o10.REPLIES = {comment_id(path, "Livia Delgado", "ADB comment (24/07/2026 PPT) asked for occupancy"): C68_REPLY}
    o10.add_replies(path)


def main(apply):
    if "--c86-reply" in sys.argv:  # user, 30 Sep: "Done." too vague here as well
        set_reply(DST, "Lucas Melo", "The text is not clear why you use 1,123", C86_REPLY)
        print("C86 reply saved", DST)
        return
    if "--zones-203" in sys.argv:  # user decision, 30 Sep
        zones_203(DST)
        print("208 -> 203 saved", DST)
        return
    if "--xref-only" in sys.argv:  # added after the first run (30 Sep)
        fix_table_xref(DST)
        print("xref saved", DST)
        return
    if "--replies-only" in sys.argv:  # table already in (first run stopped at the replies)
        replies(DST)
        print("replies saved", DST)
        return
    doc = Document(DST)
    if any(p.text.endswith(TITLE) for p in doc.paragraphs):
        raise SystemExit("already applied")
    data = add_table(doc)
    swap_figure(doc)
    for r in data:
        print("  " + " | ".join(r))
    if not apply:
        print("dry run OK - nothing saved")
        return
    doc.save(DST)
    replies(DST)
    fix_table_xref(DST)
    zones_203(DST)
    print("saved", DST)


if __name__ == "__main__":
    main("--apply" in sys.argv)
