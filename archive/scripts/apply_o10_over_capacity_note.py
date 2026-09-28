# -*- coding: utf-8 -*-
"""Explain what "over capacity" means in Output 10, IN PLACE.

Output 10 uses the phrase four times - "42 of its 53 zones over capacity", "116 are
over capacity outright" - and never says how a kerb can hold more vehicles than it has
spaces. Output 8 carries that explanation in its occupancy chapter, but Output 10 is
read on its own, and the obvious reader question goes unanswered where the number is
doing the most work: in the argument for pricing to the 85% target.

The answer is turnover. Occupancy counts every distinct vehicle that used a zone during
the hour, not how many stood in it at one moment, so a zone that turns over within the
hour serves more vehicles than it has spaces.

Two edits: the full sentence at first use, and a short clause beside the 116 in the
pricing measure, for readers who start at the measures chapter. Backup written first.
"""
import os
import re
import shutil

from docx import Document

import docx_edit as dx
import report_figures as rf

DOC = os.path.join(rf.ROOT, "Final Presentation",
                   "Output 10 - Parking Analysis Report - 25082026 (rev).docx")
BAK = DOC.replace(".docx", " (pre-over-capacity-note).docx")


def sub_in_place(par, pattern, repl):
    """Regex-replace across runs, keeping each touched run's own formatting, so text
    already accepted into black stays black."""
    m = re.search(pattern, par.text)
    if not m:
        return False
    i, j, first = m.start(), m.end(), True
    pos = 0
    for run in par.runs:
        s, e = pos, pos + len(run.text)
        pos = e
        if e <= i or s >= j:
            continue
        head = run.text[: max(0, i - s)]
        tail = run.text[max(0, j - s):] if j < e else ""
        run.text = head + (m.expand(repl) if first else "") + tail
        first = False
    return True


def main():
    shutil.copyfile(DOC, BAK)
    doc = Document(DOC)
    done, skipped = [], []

    # 1. first use, in the corridor-differences paragraph
    p = next((x for x in doc.paragraphs
              if "The Kentron commercial core runs at" in x.text), None)
    if p is not None and sub_in_place(
            p,
            r"with (\d+) of its 53 zones over capacity at their own peak",
            r"with \1 of its 53 zones over capacity at their own peak - occupancy counts "
            r"every distinct vehicle that used a zone during the hour rather than how "
            r"many stood in it at one moment, so a zone whose spaces turn over within "
            r"the hour serves more vehicles than it has spaces -"):
        done.append("explanation added at first use")
    else:
        skipped.append("first use")

    # 2. the pricing measure, where the 116 carries the argument
    p = next((x for x in doc.paragraphs
              if "are over capacity outright" in x.text), None)
    if p is not None and sub_in_place(
            p,
            r"and (\d+) are over capacity outright\.",
            r"and \1 serve more vehicles across their busiest hour than they have "
            r"spaces."):
        done.append("pricing-measure wording clarified")
    else:
        skipped.append("pricing measure")

    doc.save(DOC)
    print("backup:", os.path.basename(BAK))
    for d in done:
        print("  applied:", d)
    for s in skipped:
        print("  SKIPPED:", s)
    print("saved:", os.path.basename(DOC))


if __name__ == "__main__":
    main()
