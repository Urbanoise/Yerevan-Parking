"""Drop the "prerequisite" framing wherever it is attached to the residential-yard
opening measure, in both the Output 10 and Output 8 revisions.

The yard-opening measure may yet be dropped from the programme, so no conclusion,
table row or sequencing statement should be written as if it were a precondition
the rest of the package hangs on. The underlying measured facts are kept - the
absorptive capacity really is mostly off-street - but they are now stated as
observations rather than as a condition to be satisfied. The general
"mitigation before removal" sequencing point is kept too, restated without the
yard dependency.

Both files are edited in place (they carry hand edits made after the reviewer-comment
scripts ran, so they are not regenerated). A backup is written alongside each.
"""
import os
import shutil

from docx import Document

import report_figures as rf

PRES = os.path.join(rf.ROOT, "Final Presentation")
O10 = os.path.join(PRES, "Output 10 - Parking Analysis Report - 13082026 (rev).docx")
O8 = os.path.join(PRES, "Output 8 - Parking Surveys and Analysis Report 12082026 (rev).docx")


def span_sub(par, old, new):
    """Replace `old` with `new` inside a paragraph, even when Word has split it
    across several runs. The first touched run keeps its formatting and carries the
    replacement; the rest of the span is emptied."""
    text = par.text
    i = text.find(old)
    if i < 0:
        raise LookupError(f"not found: {old[:60]!r}")
    j = i + len(old)
    pos = 0
    first = True
    for run in par.runs:
        s, e = pos, pos + len(run.text)
        pos = e
        if e <= i or s >= j:
            continue
        head = run.text[: max(0, i - s)]
        tail = run.text[max(0, j - s):] if j < e else ""
        run.text = head + (new if first else "") + tail
        first = False
    return par


def find(doc, needle):
    for p in doc.paragraphs:
        if needle in p.text:
            return p
    raise LookupError(needle[:60])


def do_o10():
    shutil.copyfile(O10, O10.replace(".docx", " (pre-prereq-trim).docx"))
    doc = Document(O10)

    # 5.7 Phasing - the paragraph existed only to make the yards a precondition.
    # Keep the sequencing emphasis, drop the dependency and the word.
    span_sub(
        find(doc, "Because re-absorption depends on the gated off-street yards"),
        "Because re-absorption depends on the gated off-street yards (above), opening "
        "and actively managing them together with the pricing must precede kerb "
        "removal. The off-street dependency makes 'mitigation before removal' a "
        "prerequisite rather than advisory: removing kerb capacity ahead of the yards "
        "would strand the displaced demand the survey measures.",
        "This sequencing is a practical requirement rather than a counsel of "
        "perfection: removing kerb capacity before the management measures are in "
        "place would leave the displaced demand the survey measures without an "
        "organised alternative.")

    # Conclusions - "layered, not singular"
    span_sub(
        find(doc, "with the yard opening as the prerequisite the arithmetic depends on"),
        " - applied selectively by sensitivity zone, with the yard opening as the "
        "prerequisite the arithmetic depends on.",
        " - applied selectively by sensitivity zone.")

    # Conclusions - "measured, not inferred"
    span_sub(
        find(doc, "achievable but conditional on opening the gated residential"),
        "Aggregate re-absorption approaching 100% is achievable but conditional on "
        "opening the gated residential courtyards, which supply 83% of the absorptive "
        "capacity counted.",
        "Aggregate re-absorption approaching 100% is achievable, on absorptive "
        "capacity that is 83% off-street rather than kerbside.")

    # Institutional responsibility table
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if "Open residential yards (prerequisite)" in p.text:
                        span_sub(p, "Open residential yards (prerequisite)",
                                 "Open residential yards")

    doc.save(O10)
    print("O10 saved")


def do_o8():
    shutil.copyfile(O8, O8.replace(".docx", " (pre-prereq-trim).docx"))
    doc = Document(O8)

    span_sub(
        find(doc, "is therefore conditional on those yards being opened"),
        "The conclusion that displaced demand can be re-absorbed (approaching 100% in "
        "aggregate) is therefore conditional on those yards being opened and actively "
        "managed; this dependency, and the measures to address it, are carried in the "
        "companion Parking Analysis Report (Output 10).",
        "Re-absorption approaching 100% in aggregate therefore draws mainly on "
        "off-street capacity rather than on the kerb; the measures that address this "
        "are set out in the companion Parking Analysis Report (Output 10).")

    span_sub(
        find(doc, "so the courtyard dependency refers to the still-closed"),
        "so the courtyard dependency refers to the still-closed residential yards that "
        "could not be surveyed.",
        "so the off-street capacity counted above refers to the still-closed "
        "residential yards that could not be surveyed.")

    doc.save(O8)
    print("O8 saved")


if __name__ == "__main__":
    do_o10()
    do_o8()
