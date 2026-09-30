# -*- coding: utf-8 -*-
"""Output 10 - Logit comments of 24 Sep 2026: STS response copy.

Builds a NEW copy of the Logit-commented file carrying (1) the same red edits as the
clean 28092026 revision, and (2) a short threaded reply from Nikoloz Archvadze under
each of Livia Delgado's five comments. The Logit original is never modified.

The Logit copy already contains the reviewers' own untracked edits (the rewritten
delineation paragraph and the "is is" fix), so those two are asserted, not re-applied.

Check: after pass 1 the paragraph text must equal the 28092026 file exactly.

Already built on 29 Sep; apply_o10_sep_kentron53.py then corrected both files (Kentron
53 zones, survey 203) and this copy's C61 reply. A rebuild would need that script
re-run afterwards - the pass-1 check will fail against the corrected clean file.

Run:  python build_o10_logit_responses.py        (refuses to overwrite the output)
"""
import copy
import datetime as dt
import io
import os
import random
import shutil
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from apply_o10_sep_comments import (C57_CAVEAT, C85_TARIFF, C86_BODY, C86_TITLE, C90_LEAD,
                                    C61_ZONES, _clone_run_after, clone_para_after,
                                    insert_red)
from apply_o10_sep_zone_definition import DEFINITION
from docx_edit import find

SEP = "C:/Users/user/Yerevan-Parking/Final Presentation/SEP'26/"
SRC = SEP + "Output 10_Parking Analysis Report_25082026_Logit comments.docx"
CLEAN = SEP + "Output 10 - Parking Analysis Report - 28092026.docx"
DST = SEP + "Output 10_Parking Analysis Report_25082026_Logit comments - STS responses.docx"

AUTHOR, INITIALS = "Nikoloz Archvadze", "NA"
REPLIES = {  # parent comment id -> reply
    "57": "Done, caveat added.",
    "61": "Both correct: Kentron has 58 zones (53 surveyed), Komitas 53. Added a definition "
          "of zone and the counts per area.",
    "85": "Tariff rates are outside our scope and we didn't receive PCS payment data. "
          "Stated in the text.",
    "86": "Done.",
    "90": "Done, Transport Department as lead (per Output 28).",
}

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"
W16CID = "http://schemas.microsoft.com/office/word/2016/wordml/cid"


# ---------------------------------------------------------------------------
# pass 1: the red edits
# ---------------------------------------------------------------------------
def append_red_to_text(par, text):
    """Append after the last run that carries text - NOT after a comment-reference
    run, which is often the paragraph's last run in the commented file."""
    runs = [r for r in par.runs if r._r.find(qn("w:t")) is not None]
    return _clone_run_after(runs[-1], text, red=True)


def apply_edits(doc):
    append_red_to_text(find(doc, "That is why the design retains or re-establishes 1,220 spaces"),
                       C57_CAVEAT)
    insert_red(find(doc, "10,135 distinct vehicles across 208 survey zones"),
               "208 survey zones", DEFINITION + C61_ZONES)
    for needle, anchor in (("with 42 of its 53 zones over capacity at their own peak", "42 of its 53"),
                           ("with 42 of its 53 zones over capacity. Komitas", "42 of its 53"),
                           ("42 of 53 zones exceed capacity", "42 of 53")):
        insert_red(find(doc, needle), anchor, " surveyed")

    # reviewers' own edits are already in this copy
    find(doc, "Physical delineation is expected to influence driver behaviour directly")
    find(doc, "Yerevan's flat 5,000 AMD fine is modest")

    outer = find(doc, "This is not a general tariff increase")
    clone_para_after(outer, outer, C85_TARIFF)

    last_body = find(doc, "Neither this report nor the surveys undertaken establish the scale")
    cur = clone_para_after(last_body, find(doc, "Remove derelict and long-term abandoned vehicles"),
                           C86_TITLE)
    body_model = find(doc, "A kerb-management framework should include a route")
    for text in C86_BODY:
        cur = clone_para_after(cur, body_model, text)

    seq = find(doc, "This sequencing is a practical requirement")
    clone_para_after(seq, seq, C90_LEAD)


# ---------------------------------------------------------------------------
# pass 2: threaded replies (raw OOXML - python-docx has no reply API)
# ---------------------------------------------------------------------------
def _hex8(used):
    while True:
        v = f"{random.randint(0x10000000, 0x7FFFFFFE):08X}"
        if v not in used:
            used.add(v)
            return v


def add_replies(path):
    zin = zipfile.ZipFile(path)
    parts = {n: zin.read(n) for n in zin.namelist()}
    zin.close()
    parse = lambda n: etree.fromstring(parts[n])
    doc, com = parse("word/document.xml"), parse("word/comments.xml")
    ext, ids, ppl = (parse("word/commentsExtended.xml"), parse("word/commentsIds.xml"),
                     parse("word/people.xml"))

    w = lambda t: f"{{{W}}}{t}"
    used = {e.get(f"{{{W14}}}paraId") for e in com.iter(w("p"))} | \
           {e.get(f"{{{W14}}}paraId") for e in doc.iter(w("p"))}
    next_id = max(int(c.get(w("id"))) for c in com.findall(w("comment"))) + 1
    stamp = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:00Z")

    for parent_id, text in REPLIES.items():
        parent = com.xpath(f"w:comment[@w:id='{parent_id}']", namespaces={"w": W})[0]
        parent_para = parent.findall(w("p"))[-1].get(f"{{{W14}}}paraId")
        rid, para_id, durable = str(next_id), _hex8(used), _hex8(used)
        next_id += 1

        c = etree.SubElement(com, w("comment"))
        c.set(w("id"), rid); c.set(w("author"), AUTHOR); c.set(w("date"), stamp)
        c.set(w("initials"), INITIALS)
        p = etree.SubElement(c, w("p"))
        p.set(f"{{{W14}}}paraId", para_id); p.set(f"{{{W14}}}textId", "77777777")
        ppr = etree.SubElement(p, w("pPr"))
        etree.SubElement(ppr, w("pStyle")).set(w("val"), "CommentText")
        r1 = etree.SubElement(p, w("r"))
        etree.SubElement(etree.SubElement(r1, w("rPr")), w("rStyle")).set(w("val"), "CommentReference")
        etree.SubElement(r1, w("annotationRef"))
        t = etree.SubElement(etree.SubElement(p, w("r")), w("t"))
        t.text = text

        ex = etree.SubElement(ext, f"{{{W15}}}commentEx")
        ex.set(f"{{{W15}}}paraId", para_id); ex.set(f"{{{W15}}}paraIdParent", parent_para)
        ex.set(f"{{{W15}}}done", "0")
        ci = etree.SubElement(ids, f"{{{W16CID}}}commentId")
        ci.set(f"{{{W16CID}}}paraId", para_id); ci.set(f"{{{W16CID}}}durableId", durable)

        # anchor the reply on the same range as its parent
        ns = {"w": W}
        doc.xpath(f"//w:commentRangeStart[@w:id='{parent_id}']", namespaces=ns)[0].addnext(
            etree.Element(w("commentRangeStart"), {w("id"): rid}))
        doc.xpath(f"//w:commentRangeEnd[@w:id='{parent_id}']", namespaces=ns)[0].addnext(
            etree.Element(w("commentRangeEnd"), {w("id"): rid}))
        ref_run = doc.xpath(f"//w:r[w:commentReference[@w:id='{parent_id}']]", namespaces=ns)[0]
        nr = etree.Element(w("r"))
        etree.SubElement(etree.SubElement(nr, w("rPr")), w("rStyle")).set(w("val"), "CommentReference")
        etree.SubElement(nr, w("commentReference")).set(w("id"), rid)
        ref_run.addnext(nr)

    if not ppl.xpath(f"w15:person[@w15:author='{AUTHOR}']", namespaces={"w15": W15}):
        person = etree.SubElement(ppl, f"{{{W15}}}person")
        person.set(f"{{{W15}}}author", AUTHOR)
        pi = etree.SubElement(person, f"{{{W15}}}presenceInfo")
        pi.set(f"{{{W15}}}providerId", "None"); pi.set(f"{{{W15}}}userId", AUTHOR)

    ser = lambda x: etree.tostring(x, xml_declaration=True, encoding="UTF-8", standalone=True)
    parts.update({"word/document.xml": ser(doc), "word/comments.xml": ser(com),
                  "word/commentsExtended.xml": ser(ext), "word/commentsIds.xml": ser(ids),
                  "word/people.xml": ser(ppl)})
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in parts.items():
            zout.writestr(name, data)
    with open(path, "wb") as fh:
        fh.write(buf.getvalue())


def main():
    if os.path.exists(DST):
        raise SystemExit(f"refusing to overwrite {DST}")
    doc = Document(SRC)
    apply_edits(doc)
    got = [p.text for p in doc.paragraphs]
    want = [p.text for p in Document(CLEAN).paragraphs]
    if got != want:
        diffs = [i for i, (a, b) in enumerate(zip(got, want)) if a != b]
        raise SystemExit(f"text differs from the clean revision: {len(got)} vs {len(want)} "
                         f"paragraphs, first differences at {diffs[:5]}")
    doc.save(DST)
    add_replies(DST)
    print("wrote", DST)


if __name__ == "__main__":
    main()
