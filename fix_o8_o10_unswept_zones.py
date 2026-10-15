# -*- coding: utf-8 -*-
"""OCT'26 copies of Output 8 / Output 10: drop the five unswept Kentron zones.

Kentron K02–K05 and K18 (52 spaces) were drawn for supply but never swept for
occupancy. Since 15 Oct 2026 compute_field_survey_metrics.mjs leaves them out of the
survey layer, so every survey-area figure rests on the same 203 zones / 2,878 spaces:
Kentron 586 -> 534 spaces and 408 -> 356 removed, six areas 2,930 -> 2,878 and
1,886 -> 1,834 removed. The sentence explaining the 52-space gap goes. Corridor totals
(4,843 removed, 1,252 retained) are unaffected.

Edits the saved files in place (they carry hand edits). Only the changed words are
red; a deletion is simply removed. Figures 14, 16 and 20 are swapped for re-renders.

Usage: python fix_o8_o10_unswept_zones.py
"""
import copy
import os

from docx import Document
from docx.text.run import Run

import docx_edit as dx
import gen_o8_report_figures as g
import report_figures as rf

OCT = os.path.join(rf.ROOT, "Final Presentation", "OCT'26")
O8 = os.path.join(OCT, "Output 8 - Parking Surveys and Analysis Report 15102026 (rev).docx")
O10 = os.path.join(OCT, "Output 10 - Parking Analysis Report 15102026 (rev).docx")
MAPS = os.path.join(OCT, "O8 app maps")
g.OUT = os.path.join(OCT, "charts")


def replace_span(par, old, new):
    """Replace the first `old` in a paragraph without touching the rest of it: the
    characters of `old` are cut from whichever runs hold them, and `new` (if any) goes
    in as one red run at that spot, formatted like the run it lands in."""
    runs = par.runs
    text = "".join(r.text for r in runs)
    start = text.find(old)
    if start < 0:
        raise ValueError(f"not found: {old!r}")
    end = start + len(old)
    pos, first = 0, True
    for run in runs:
        r0, r1 = pos, pos + len(run.text)
        pos = r1
        if r1 <= start or r0 >= end:
            continue
        t = run.text
        a, b = max(start, r0) - r0, min(end, r1) - r0
        if first:
            first = False
            if new:
                ins = copy.deepcopy(run._r)
                run._r.addnext(ins)
                rin = Run(ins, run._parent)
                rin.text = new
                rin.font.color.rgb = dx.RED
                if t[b:]:
                    tail = copy.deepcopy(run._r)
                    ins.addnext(tail)
                    Run(tail, run._parent).text = t[b:]
                run.text = t[:a]
                continue
        run.text = t[:a] + t[b:]
    after = "".join(r.text for r in par.runs)
    assert after == text[:start] + new + text[end:], (old, after[:120])
    return par


def cell_with(table, text):
    for row in table.rows:
        for c in row.cells:
            if c.text.strip() == text:
                return c
    raise ValueError(f"cell not found: {text!r}")


def output8(f):
    k, al = f["areas"]["kentron"], f["all"]
    assert (k["spaces"], k["removed"], al["removed"], al["displaced"]) == (534, 356, 1834, 673)
    empty = al["removed"] - al["displaced"]                     # 1,161
    doc = Document(O8)

    replace_span(dx.find(doc, "In the surveyed areas the design removes 1,886 spaces"),
                 "1,886", f"{al['removed']:,}")

    p = dx.find(doc, "five Kentron zones were inventoried on the supply side")
    replace_span(p, "36% of the 2,878 spaces in the zones where occupancy was recorded.",
                 "36% of their 2,878 spaces.")
    replace_span(p, " That is 52 fewer than the 2,930 spaces in Table 9, because five Kentron "
                    "zones were inventoried on the supply side but not swept for occupancy; "
                    "they are excluded from every occupancy figure in this chapter, which is "
                    "why Kentron is measured against 534 spaces rather than 586.", "")

    p = dx.find(doc, "The conceptual design removes 1,886 spaces in the six surveyed areas")
    replace_span(p, "1,886", f"{al['removed']:,}")
    replace_span(p, "1,213", f"{empty:,}")
    replace_span(p, "about 64%", f"about {round(empty * 100 / al['removed'])}%")

    t9 = dx.find_table(doc, "Survey date", "Formal spaces")
    dx.set_cell(cell_with(t9, "586"), "534")
    dx.set_cell(cell_with(t9, "2,930"), "2,878")
    t11 = dx.find_table(doc, "Spaces removed", "Displaced at peak")
    dx.set_cell(cell_with(t11, "408"), "356")
    dx.set_cell(cell_with(t11, "1,886"), "1,834")

    for needle, path in {
        "Location of the six field survey areas": os.path.join(MAPS, "areas.jpg"),
        "occupancy by zone in the six survey areas": os.path.join(MAPS, "occ_grid.jpg"),
        "Parking removed and demand displaced, by area": g.fig16_displacement(f),
    }.items():
        dx.swap_figure_above(doc, needle, path)

    for needle in ("586", "2,930", "408", "1,886", "1,213", "52 fewer"):
        hits = [p.text[:80] for p in dx.find_all(doc, needle)]
        hits += [c.text[:80] for t in doc.tables for r in t.rows for c in r.cells if needle in c.text]
        if hits:
            print(f"O8 CHECK {needle!r}:", hits)
    doc.save(O8)
    print("saved", O8)


def output10(f):
    doc = Document(O10)
    for needle in ("in the 1,886 spaces to be removed", "the design removes 1,886 spaces"):
        replace_span(dx.find(doc, needle), "1,886", f"{f['all']['removed']:,}")
    for needle in ("586", "2,930", "408", "1,886"):
        if dx.find_all(doc, needle):
            print(f"O10 CHECK {needle!r}")
    doc.save(O10)
    print("saved", O10)


if __name__ == "__main__":
    f = rf.survey()
    output8(f)
    output10(f)
