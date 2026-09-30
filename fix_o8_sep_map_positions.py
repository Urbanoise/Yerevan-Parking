# -*- coding: utf-8 -*-
"""Output 8 - Sep 2026 response copy: move four maps out of the List of Figures.

apply_o8_sep_maps.py (first run, 30 Sep) anchored four figures on caption text with an
unstyled find(), which matched the List of Figures entries at the front of the report.
The picture + caption pairs for the regulatory-zone, configuration, signage and
marking maps therefore sat inside the list. This moves each pair to directly after
the real caption it belongs under. The script itself is now fixed.

Run:  python fix_o8_sep_map_positions.py            (dry run)
      python fix_o8_sep_map_positions.py --apply
"""
import sys

from docx import Document

from apply_o8_sep_comments import DST
from docx_edit import find_styled

MOVES = {  # misplaced caption -> real caption it belongs under
    "Location of on-corridor on-street parking by regulatory zone":
        "Figure 1 - On-corridor on-street supply by regulatory zone",
    "Location of on-corridor on-street parking by configuration":
        "Figure 2 - On-corridor on-street supply by marked configuration",
    "Location of parking signage": "Figure 4 - Parking signage coverage",
    "Location of pavement marking": "Figure 5 - Pavement marking coverage",
}


def main(apply):
    doc = Document(DST)
    paras = doc.paragraphs
    body_start = next(i for i, p in enumerate(paras)
                      if p.style.name == "Heading 1" and p.text.strip() == "Executive Summary")
    for title, target in MOVES.items():
        hits = [i for i, p in enumerate(paras[:body_start])
                if p.style.name == "Caption" and p.text.endswith(title)]
        if len(hits) != 1:
            raise SystemExit(f"{len(hits)} misplaced captions for {title!r}")
        cap = paras[hits[0]]
        pic = paras[hits[0] - 1]
        if "graphic" not in pic._p.xml:
            raise SystemExit(f"no picture before {title!r}")
        anchor = find_styled(doc, target, "Caption")
        anchor._p.addnext(pic._p)
        pic._p.addnext(cap._p)
        paras = doc.paragraphs

    front = [p.text for p in doc.paragraphs[:body_start] if p.style.name == "Caption"]
    if front:
        raise SystemExit(f"captions still in the front matter: {front}")
    if apply:
        doc.save(DST)
        print("saved", DST)
    else:
        print("dry run OK - nothing saved")


if __name__ == "__main__":
    main("--apply" in sys.argv)
